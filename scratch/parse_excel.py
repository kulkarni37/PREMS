import openpyxl

def parse_past_projects(excel_path='D:/projectlist 5 year.xlsx'):
    wb = openpyxl.load_workbook(excel_path)
    all_projects = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_count = 0
        for r in range(2, ws.max_row + 1):
            c1 = ws.cell(row=r, column=1).value
            c2 = ws.cell(row=r, column=2).value
            c3 = ws.cell(row=r, column=3).value if ws.max_column >= 3 else None

            title = None
            group_no = None
            guide = None

            if sheet_name == '2018-22':
                if c2 and str(c2).strip():
                    title = str(c2).strip()
                elif c1 and str(c1).strip() and not str(c1).strip().isdigit():
                    title = str(c1).strip()
                
                if c1 and str(c1).strip():
                    c1_str = str(c1).strip()
                    if c1_str.isdigit():
                        group_no = f'G{c1_str}'
                    elif c1_str.startswith('G'):
                        group_no = c1_str

            elif sheet_name == '2019-23':
                if c1 and str(c1).strip():
                    group_no = str(c1).strip()
                if c2 and str(c2).strip():
                    title = str(c2).strip()
                if c3 and str(c3).strip():
                    guide = ' '.join(str(c3).strip().split())

            else: # 2020-24, 2021-25, 2022-25
                if c1 and str(c1).strip():
                    group_no = str(c1).strip()
                if c2 and str(c2).strip():
                    title = str(c2).strip()

            if title:
                clean_title = ' '.join(str(title).split()).strip('\"\'')
                if clean_title and clean_title.lower() not in ['title of project', 'project title', 'slno', 'group no.', 'guide']:
                    all_projects.append({
                        'title': clean_title,
                        'source_period': sheet_name,
                        'group_no': group_no,
                        'guide_name': guide
                    })
                    sheet_count += 1

        print(f"Sheet {sheet_name}: {sheet_count} valid projects")

    print(f"Total projects parsed: {len(all_projects)}")
    return all_projects

if __name__ == '__main__':
    projects = parse_past_projects()
    print("\nSample records:")
    for p in projects[:10]:
        print(p)
